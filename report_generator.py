import csv
from openpyxl import Workbook
from openpyxl.styles import Font

#createWorkbook
wb2=Workbook()
ws=wb2.active
#open previous csv file
with open('books.csv',mode='r') as csvfile:
 reader=csv.reader(csvfile)
 for row in reader:
  ws.append(row)
#bold header
ft=Font(bold=True)
for row in ws['A1:C1']:
    for cell in row:
        cell.font=ft
#count the books
totalbooks=0
for row in ws:
   totalbooks=totalbooks+1
totalbooks=totalbooks-1
emptyrow=ws.max_row+1
ws.cell(row=emptyrow,column=1,value=totalbooks)
#adjusting col's width
for col in ws.columns:
   max_length=0
   for cell in col:
      data=str(cell.value)  
      if (len(data))>max_length:
         max_length=len(data)
   firstcol=col[0].column_letter
   ws.column_dimensions[firstcol].width=max_length+2
wb2.save('report.xlsx')



      
    

